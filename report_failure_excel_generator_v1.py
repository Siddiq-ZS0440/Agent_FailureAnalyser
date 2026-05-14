from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError as exc:
    raise SystemExit("Missing dependency: openpyxl. Install it using: pip install openpyxl") from exc

LogCallback = Optional[Callable[[str], None]]
SUMMARY_SHEET = "Summary"
DETAILS_SHEET = "Categorized Failures"


def _emit(logger: LogCallback, message: str) -> None:
    if logger:
        logger(message)


def load_json(json_path: Path) -> Dict[str, Any]:
    with json_path.open("r", encoding="utf-8") as file:
        return json.load(file)


def safe_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def autosize_columns(ws, max_width: int = 80) -> None:
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = 0
        for cell in column_cells:
            value = safe_text(cell.value)
            if len(value) > max_length:
                max_length = len(value)
        ws.column_dimensions[column_letter].width = min(max_length + 2, max_width)


def style_header_row(ws, row_number: int) -> None:
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    for cell in ws[row_number]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border


def style_body(ws) -> None:
    thin_border = Border(
        left=Side(style="thin", color="D9E2F3"),
        right=Side(style="thin", color="D9E2F3"),
        top=Side(style="thin", color="D9E2F3"),
        bottom=Side(style="thin", color="D9E2F3"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = thin_border


def build_summary_sheet(wb: Workbook, data: Dict[str, Any]) -> None:
    ws = wb.active
    ws.title = SUMMARY_SHEET

    ws.append(["Field", "Value"])
    ws.append(["Source HTML File", safe_text(data.get("source_html_file"))])
    ws.append(["Source HTML Path", safe_text(data.get("source_html_path"))])
    ws.append(["Report Type", safe_text(data.get("report_type"))])
    ws.append(["Report Name", safe_text(data.get("report_name"))])
    ws.append(["Total Test Count", data.get("total_test_count", 0)])
    ws.append(["Failed Test Count", data.get("failed_test_count", 0)])
    ws.append([])
    ws.append(["Failure Category", "Count"])

    category_summary = data.get("failure_category_summary", {})
    if isinstance(category_summary, dict):
        for category, count in category_summary.items():
            ws.append([safe_text(category), count])

    style_header_row(ws, 1)
    style_header_row(ws, 9)
    style_body(ws)
    autosize_columns(ws)
    ws.freeze_panes = "A2"


def build_details_sheet(wb: Workbook, data: Dict[str, Any]) -> None:
    ws = wb.create_sheet(DETAILS_SHEET)

    headers = [
        "S.No",
        "Test Name",
        "Status",
        "Failure Category",
        "Failure Message",
        "Stack Trace",
    ]
    ws.append(headers)

    failed_tests: List[Dict[str, Any]] = data.get("categorized_failed_test_list", [])
    for index, test in enumerate(failed_tests, start=1):
        ws.append(
            [
                index,
                safe_text(test.get("test_name")),
                safe_text(test.get("status")),
                safe_text(test.get("failure_category")),
                safe_text(test.get("failure_message")),
                safe_text(test.get("stack_trace")),
            ]
        )

    style_header_row(ws, 1)
    style_body(ws)
    autosize_columns(ws)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 45
    ws.column_dimensions["E"].width = 70
    ws.column_dimensions["F"].width = 80
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def generate_excel(
    input_json_path: str,
    output_excel_path: str = "",
    logger: LogCallback = None,
) -> Path:
    input_path = Path(input_json_path).resolve()
    _emit(logger, f"Reading categorization JSON: {input_path}")
    data = load_json(input_path)

    output_path = Path(output_excel_path).resolve() if output_excel_path else input_path.parent / "report_failure_categorization.xlsx"

    wb = Workbook()
    build_summary_sheet(wb, data)
    build_details_sheet(wb, data)
    wb.save(output_path)
    _emit(logger, f"Saved Excel workbook to {output_path}")
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate Excel report from report_failure_categorization.json.")
    parser.add_argument("categorization_json", help="Path to report_failure_categorization.json")
    parser.add_argument("--output", default="", help="Optional output Excel path. Defaults to report_failure_categorization.xlsx in the same folder.")
    args = parser.parse_args()

    output_path = generate_excel(args.categorization_json, args.output, logger=print)
    print(str(output_path.resolve()))


if __name__ == "__main__":
    main()
